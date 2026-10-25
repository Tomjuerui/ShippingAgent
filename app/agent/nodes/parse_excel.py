"""Excel 解析（FR-04）：多工作表、表头偏移、合并单元格展开。"""

from __future__ import annotations

import datetime as dt
from typing import Any

from app.agent.nodes.unpack_eml import Payload
from app.extraction.schemas import UIRBlock, UIRDocument, UIRTable

_HEADER_KEYWORDS: dict[str, tuple[str, ...]] = {
    "charge": ("charge", "description", "item", "particular", "fee", "费用", "项目", "品名"),
    "amount": ("amount", "金额", "总额", "total", "小计", "subtotal"),
    "currency": ("currency", "ccy", "curr", "币种"),
    "quantity": ("qty", "quantity", "数量"),
    "unit_price": ("unit price", "unitprice", "rate", "单价"),
    "payment": ("payment", "prepaid", "collect", "预付", "到付", "term"),
}
_HEADER_HIT = {key: tuple(k.lower() for k in vals) for key, vals in _HEADER_KEYWORDS.items()}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        text = f"{value:.10f}".rstrip("0").rstrip(".")
        return text or "0"
    return str(value).strip()


def _header_score(row: list[str]) -> int:
    joined = " ".join(row).lower()
    hits = 0
    for keywords in _HEADER_HIT.values():
        if any(keyword in joined for keyword in keywords):
            hits += 1
    return hits


def _detect_header(matrix: list[list[str]], scan_rows: int = 30) -> int | None:
    best_index = None
    best_score = 0
    for index, row in enumerate(matrix[:scan_rows]):
        score = _header_score(row)
        if score > best_score:
            best_index, best_score = index, score
    if best_score >= 2:
        return best_index
    return None


def _matrix_from_xlsx(data: bytes) -> list[tuple[str, list[list[str]], dict]]:
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    sheets: list[tuple[str, list[list[str]], dict]] = []
    for ws in workbook.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        if max_row == 0 or max_col == 0:
            continue
        values: dict[tuple[int, int], Any] = {}
        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            for cell in row:
                values[(cell.row, cell.column)] = cell.value
        for merged in ws.merged_cells.ranges:
            top_left = values.get((merged.min_row, merged.min_col))
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    if (r, c) != (merged.min_row, merged.min_col):
                        values[(r, c)] = top_left
        matrix = [
            [_cell_text(values.get((r, c))) for c in range(1, max_col + 1)]
            for r in range(1, max_row + 1)
        ]
        sheets.append((ws.title, matrix, {"merged_ranges": len(ws.merged_cells.ranges)}))
    return sheets


def _matrix_from_xls(data: bytes) -> list[tuple[str, list[list[str]], dict]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheets: list[tuple[str, list[list[str]], dict]] = []
    for sheet in book.sheets():
        values: dict[tuple[int, int], Any] = {}
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                values[(r, c)] = sheet.cell_value(r, c)
        merged = list(getattr(sheet, "merged_cells", []) or [])
        for rlo, rhi, clo, chi in merged:
            top_left = values.get((rlo, clo))
            for r in range(rlo, rhi):
                for c in range(clo, chi):
                    if (r, c) != (rlo, clo):
                        values[(r, c)] = top_left
        matrix = [
            [_cell_text(values.get((r, c))) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        sheets.append((sheet.name, matrix, {"merged_ranges": len(merged)}))
    return sheets


def _parse_sheet(
    sheet_name: str, matrix: list[list[str]], extra: dict, uir_id: str
) -> tuple[list[UIRBlock], list[str]]:
    blocks: list[UIRBlock] = []
    warnings: list[str] = []
    matrix = [row for row in matrix if any(cell for cell in row)]
    if not matrix:
        return blocks, warnings
    header_index = _detect_header(matrix)
    if header_index is None:
        header_index = 0
        warnings.append(f"工作表 {sheet_name} 未识别到标准表头，按首行表头尽力解析")
    header = [cell for cell in matrix[header_index]]
    pre_text = "\n".join(
        " | ".join(cell for cell in row if cell) for row in matrix[:header_index]
    ).strip()
    if pre_text:
        blocks.append(
            UIRBlock(
                block_id=f"{uir_id}-{sheet_name}-pre",
                kind="text",
                text=pre_text,
                locator={"sheet": sheet_name, "row_range": [1, header_index]},
            )
        )
    rows = [
        row
        for row in matrix[header_index + 1 :]
        if any(cell.strip() for cell in row)
    ]
    if not rows:
        warnings.append(f"工作表 {sheet_name} 表头之后没有明细行")
    blocks.append(
        UIRBlock(
            block_id=f"{uir_id}-{sheet_name}-table",
            kind="table",
            table=UIRTable(
                table_id=f"{uir_id}-{sheet_name}",
                headers=[cell for cell in header],
                rows=rows,
                locator={
                    "sheet": sheet_name,
                    "header_row": header_index + 1,
                    "row_count": len(rows),
                    **extra,
                },
            ),
            locator={"sheet": sheet_name, "header_row": header_index + 1},
        )
    )
    return blocks, warnings


def parse_excel(payload: Payload, uir_id: str | None = None) -> UIRDocument:
    uir_id = uir_id or f"uir-{payload.sha256[:10]}-excel"
    if payload.file_type == "xlsx":
        sheets = _matrix_from_xlsx(payload.data)
    elif payload.file_type == "xls":
        sheets = _matrix_from_xls(payload.data)
    else:
        raise ValueError(f"不支持的 Excel 类型：{payload.file_type}")
    blocks: list[UIRBlock] = []
    warnings: list[str] = []
    for sheet_name, matrix, extra in sheets:
        sheet_blocks, sheet_warnings = _parse_sheet(sheet_name, matrix, extra, uir_id)
        blocks.extend(sheet_blocks)
        warnings.extend(sheet_warnings)
    return UIRDocument(
        uir_id=uir_id,
        source=payload.source,
        content_node_id=payload.node_id,
        blocks=blocks,
        parse_mode="excel",
        parse_quality="high" if any(b.kind == "table" for b in blocks) else "low",
        warnings=warnings,
    )
