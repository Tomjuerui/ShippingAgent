"""生成 golden 评测夹具与 cases.json（可重复执行）。

用法：python eval/make_fixtures.py
产物：eval/golden/files/*、eval/golden/bills/*、eval/golden/cases.json
标注规范见 eval/golden/README.md。
"""

from __future__ import annotations

import json
from email.message import EmailMessage
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent / "golden"
FILES = ROOT / "files"
BILLS = ROOT / "bills"
FILES.mkdir(parents=True, exist_ok=True)
BILLS.mkdir(parents=True, exist_ok=True)

PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def write_xlsx(path: Path, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "BILL"
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def write_pdf(path: Path, text: str) -> Path:
    def esc(value: str) -> str:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = f"BT /F1 10 Tf 40 750 Td ({esc(text)}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    header = b"%PDF-1.4\n"
    body = b""
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(header) + len(body))
        body += f"{index} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref_offset = len(header) + len(body)
    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n" + "".join(
        f"{offset:010d} 00000 n \n" for offset in offsets
    )
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    )
    path.write_bytes(header + body + xref.encode("latin-1") + trailer.encode("latin-1"))
    return path


def write_eml(path: Path, attachment: Path) -> Path:
    message = EmailMessage()
    message["Subject"] = "Debit Note DN-2026-2004"
    message["From"] = "billing@destination-agent.com"
    message["To"] = "ap@origin-agent.com"
    message["Message-ID"] = "<dn-2026-2004@destination-agent.com>"
    message["Date"] = "Sat, 26 Sep 2026 09:00:00 +0800"
    message.set_content("Please see the attached debit note.")
    message.add_attachment(
        attachment.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=attachment.name,
    )
    path.write_bytes(bytes(message))
    return path


def write_json(path: Path, payload: dict) -> Path:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def charge(name, qty, unit, amount, currency, term, snippet) -> dict:
    return {
        "line_id": None,
        "original_charge_name": name,
        "quantity": qty,
        "unit_price": unit,
        "amount": amount,
        "currency": currency,
        "payment_term": term,
        "snippet": snippet,
        "locator": None,
        "confidence": {
            "original_charge_name": 0.99,
            "quantity": 0.99 if qty is not None else 0.0,
            "unit_price": 0.99 if unit is not None else 0.0,
            "amount": 0.99,
            "currency": 0.99,
            "payment_term": 0.99 if term else 0.0,
        },
    }


def total(kind, amount, currency) -> dict:
    return {"kind": kind, "amount": amount, "currency": currency}


def extraction(
    doc_type,
    ref,
    date,
    issuer,
    recipient,
    charges,
    totals,
    bl=None,
    container=None,
) -> dict:
    return {
        "extraction_id": None,
        "document_type": doc_type,
        "reference_no": ref,
        "document_date": date,
        "shipment_info": {
            "bl_no": bl,
            "booking_no": None,
            "container_no": container,
            "vessel_voyage": None,
            "pol": None,
            "pod": None,
        },
        "issuer": issuer,
        "recipient": recipient,
        "charges": charges,
        "totals": totals,
        "direction_statement": None,
        "field_confidence": {
            "document_type": 0.99,
            "reference_no": 0.99 if ref else 0.5,
            "document_date": 0.99,
            "issuer": 0.99,
            "recipient": 0.99,
            "bl_no": 0.99 if bl else 0.5,
            "container_no": 0.99 if container else 0.5,
        },
    }


def expected(
    doc_type,
    ref,
    date,
    issuer,
    recipient,
    original,
    standard,
    amount,
    currency,
    term,
    posting,
    operation,
    match_level,
    review_reason=None,
    changed_fields=None,
) -> dict:
    return {
        "document_type": doc_type,
        "reference_no": ref,
        "document_date": date,
        "issuer": issuer,
        "recipient": recipient,
        "original_charge_name": original,
        "charge_item": standard,
        "amount": amount,
        "currency": currency,
        "payment_term": term,
        "posting_type": posting,
        "operation": operation,
        "match_level": match_level,
        "review_reason": review_reason,
        "changed_fields_contains": changed_fields or [],
    }


HEADER = ["Charge Description", "Qty", "Unit Price", "Amount", "Ccy", "Term"]


def build_files() -> None:
    # 1 Excel 标准 DN
    write_xlsx(
        FILES / "excel_standard.xlsx",
        [
            ["DEBIT NOTE DN-2026-2001"],
            ["Date: 2026-09-20", "Issuer: Destination Agent", "Recipient: Origin Agent"],
            HEADER,
            ["H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID"],
            ["Total", None, None, 40.0, "USD", None],
        ],
    )

    # 2 Word Credit Note 负数
    document = Document()
    document.add_paragraph("CREDIT NOTE CN-2026-2002")
    document.add_paragraph("Date: 2026-09-22")
    document.add_paragraph("Issuer: Origin Agent")
    document.add_paragraph("Recipient: Destination Agent")
    table = document.add_table(rows=3, cols=5)
    for index, value in enumerate(["Description", "Qty", "Unit Price", "Amount", "Currency"]):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(["Doc Fee", "1", "35", "-35.00", "USD"]):
        table.rows[1].cells[index].text = value
    for index, value in enumerate(["Total", "", "", "-35.00", "USD"]):
        table.rows[2].cells[index].text = value
    document.save(FILES / "word_credit_note.docx")

    # 3 文本 PDF（更新）
    write_pdf(
        FILES / "pdf_update.pdf",
        "DEBIT NOTE DN-2026-2003 DATE 2026-09-25 ISSUER Destination Agent "
        "RECIPIENT Origin Agent MBL OOLU4055331281 CNTR TIIU8975975 "
        "THC 1 100 100 USD CC TOTAL 100.00 USD",
    )

    # 4 EML + Excel 附件
    attachment = write_xlsx(
        FILES / "nested_bill.xlsx",
        [
            ["DEBIT NOTE DN-2026-2004"],
            ["Date: 2026-09-26", "Issuer: Destination Agent", "Recipient: Origin Agent"],
            HEADER,
            ["Trucking Fee", 1, 50.0, 50.0, "USD", "PREPAID"],
            ["Total", None, None, 50.0, "USD", None],
        ],
    )
    write_eml(FILES / "eml_with_attachment.eml", attachment)

    # 5 图片（OCR 录制回放）
    (FILES / "image_scan.png").write_bytes(PNG)

    # 6 最小格式账单 L2 派生键（无单据号输入）
    write_xlsx(
        FILES / "minimal_l2.xlsx",
        [
            ["DEBIT NOTE"],
            ["Date: 2026-09-27", "Issuer: Destination Agent", "Recipient: Origin Agent"],
            ["MBL NO.: OOLU4055331281"],
            HEADER,
            ["THC", 1, 100.0, 100.0, "USD", "CC"],
            ["Total", None, None, 100.0, "USD", None],
        ],
    )

    # 7 多币种
    write_xlsx(
        FILES / "multi_currency.xlsx",
        [
            ["DEBIT NOTE DN-2026-2007"],
            ["Date: 2026-09-28", "Issuer: Destination Agent", "Recipient: Origin Agent"],
            HEADER,
            ["O/F", 1, 100.0, 100.0, "USD", "PP"],
            ["THC", 1, 200.0, 200.0, "CNY", "CC"],
            ["Total", None, None, 100.0, "USD", None],
            ["Total", None, None, 200.0, "CNY", None],
        ],
    )

    # 8 勾稽不平
    write_xlsx(
        FILES / "sum_mismatch.xlsx",
        [
            ["DEBIT NOTE DN-2026-2008"],
            ["Date: 2026-09-29", "Issuer: Destination Agent", "Recipient: Origin Agent"],
            HEADER,
            ["H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID"],
            ["Total", None, None, 99.0, "USD", None],
        ],
    )

    # 9 主体无法对应
    write_xlsx(
        FILES / "entity_unresolved.xlsx",
        [
            ["DEBIT NOTE DN-2026-2009"],
            ["Date: 2026-09-29", "Issuer: Unknown Co.", "Recipient: Other Co."],
            HEADER,
            ["H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID"],
            ["Total", None, None, 40.0, "USD", None],
        ],
    )

    # 10 数据冲突
    write_xlsx(
        FILES / "data_conflict.xlsx",
        [
            ["DEBIT NOTE DN-2026-2010"],
            ["Date: 2026-09-30", "Issuer: Origin Agent", "Recipient: Another Co."],
            HEADER,
            ["H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID"],
            ["Total", None, None, 40.0, "USD", None],
        ],
    )

    # 11 匹配歧义
    write_xlsx(
        FILES / "ambiguous_match.xlsx",
        [
            ["DEBIT NOTE DN-2026-2011"],
            ["Date: 2026-10-01", "Issuer: Destination Agent", "Recipient: Origin Agent"],
            HEADER,
            ["H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID"],
            ["Total", None, None, 40.0, "USD", None],
        ],
    )


def extended_bill(bill_id: str, ref: str, date: str, amount: float, unit: float) -> dict:
    return {
        "bill_id": bill_id,
        "version": 1,
        "document_info": {
            "document_type": "Debit Note",
            "reference_no": ref,
            "document_date": date,
        },
        "shipment_info": {
            "bl_no": "OOLU4055331281",
            "booking_no": "BK-2026-889",
            "container_no": "TIIU8975975",
        },
        "transaction_party": {"issuer": "Destination Agent", "recipient": "Origin Agent"},
        "charges": [
            {
                "charge_item": "Terminal Handling Charge"
                if "THC" in ref or bill_id == "BILL-2003"
                else "Handling Charge",
                "original_charge_name": "THC" if bill_id == "BILL-2003" else "H/D Fee",
                "quantity": 1 if bill_id == "BILL-2003" else 2,
                "unit_price": unit,
                "amount": amount,
                "currency": "USD",
                "payment_term": "CC" if bill_id == "BILL-2003" else "PP",
            }
        ],
        "financial_attributes": {"accounting_entity": "Origin Agent", "posting_type": "AP"},
    }


def build_bills() -> None:
    write_json(
        BILLS / "BILL-2003.json",
        extended_bill("BILL-2003", "DN-2026-2003", "2026-09-25", 90.0, 90.0),
    )
    write_json(
        BILLS / "BILL-MIN-2006.json",
        {
            "bill_id": "BILL-MIN-2006",
            "document_info": {
                "document_type": "Debit Note",
                "reference_no": "OOLU4055331281",
                "document_date": "2026-09-27",
            },
            "transaction_party": {
                "issuer": "Destination Agent",
                "recipient": "Origin Agent",
            },
            "charge_details": {
                "original_charge_name": "THC",
                "charge_item": "Terminal Handling Charge",
                "quantity": 1,
                "unit_price": 100.0,
                "amount": 100.0,
                "currency": "USD",
                "payment_term": "CC",
            },
            "audit_trail": {"ai_reasoning": ["历史最小格式账单"]},
        },
    )
    conflict = extended_bill("BILL-2010", "DN-2026-2010", "2026-09-30", 40.0, 20.0)
    conflict["transaction_party"] = {
        "issuer": "Destination Agent",
        "recipient": "Origin Agent",
    }
    write_json(BILLS / "BILL-2010.json", conflict)
    for index in (1, 2):
        write_json(
            BILLS / f"BILL-2011-{index}.json",
            extended_bill(f"BILL-2011-{index}", "DN-2026-2011", "2026-10-01", 40.0, 20.0),
        )


def build_cases() -> None:
    cases = []

    cases.append(
        {
            "id": "excel_standard",
            "file": "files/excel_standard.xlsx",
            "bills": [],
            "match_keywords": ["H/D Fee", "DN-2026-2001"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2001",
                "2026-09-20",
                "Destination Agent",
                "Origin Agent",
                [charge("H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID", "H/D Fee 2 20 40 USD PREPAID")],
                [total("total", 40.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2001", "2026-09-20", "Destination Agent",
                    "Origin Agent", "H/D Fee", "Handling Charge", 40.0, "USD", "PP",
                    "AP", "CREATE", "NONE",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "word_credit_note",
            "file": "files/word_credit_note.docx",
            "bills": [],
            "match_keywords": ["Doc Fee", "CN-2026-2002"],
            "ocr": None,
            "extraction": extraction(
                "CREDIT NOTE",
                "CN-2026-2002",
                "2026-09-22",
                "Origin Agent",
                "Destination Agent",
                [charge("Doc Fee", 1, 35.0, -35.0, "USD", None, "Doc Fee 1 35 -35 USD")],
                [total("total", -35.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Credit Note", "CN-2026-2002", "2026-09-22", "Origin Agent",
                    "Destination Agent", "Doc Fee", "Documentation Fee", -35.0, "USD", None,
                    "AP", "CREATE", "NONE",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "pdf_update",
            "file": "files/pdf_update.pdf",
            "bills": ["bills/BILL-2003.json"],
            "match_keywords": ["THC", "DN-2026-2003"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2003",
                "2026-09-25",
                "Destination Agent",
                "Origin Agent",
                [charge("THC", 1, 100.0, 100.0, "USD", "CC", "THC 1 100 100 USD CC")],
                [total("total", 100.0, "USD")],
                bl="OOLU4055331281",
                container="TIIU8975975",
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2003", "2026-09-25", "Destination Agent",
                    "Origin Agent", "THC", "Terminal Handling Charge", 100.0, "USD", "CC",
                    "AP", "UPDATE", "L1", None, ["charge_details.amount"],
                )
            ],
        }
    )

    cases.append(
        {
            "id": "eml_attachment",
            "file": "files/eml_with_attachment.eml",
            "bills": [],
            "match_keywords": ["Trucking Fee", "DN-2026-2004"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2004",
                "2026-09-26",
                "Destination Agent",
                "Origin Agent",
                [charge("Trucking Fee", 1, 50.0, 50.0, "USD", "PREPAID", "Trucking Fee 1 50 50 USD PREPAID")],
                [total("total", 50.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2004", "2026-09-26", "Destination Agent",
                    "Origin Agent", "Trucking Fee", "Trucking Charge", 50.0, "USD", "PP",
                    "AP", "CREATE", "NONE",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "image_ocr",
            "file": "files/image_scan.png",
            "bills": [],
            "match_keywords": ["Customs Fee", "DN-2026-2005"],
            "ocr": {
                "lines": [
                    "DEBIT NOTE DN-2026-2005",
                    "Date 2026-09-26 Issuer Destination Agent Recipient Origin Agent",
                    "Customs Fee 80.00 USD PREPAID",
                    "Total 80.00 USD",
                ],
                "confidence": 0.96,
            },
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2005",
                "2026-09-26",
                "Destination Agent",
                "Origin Agent",
                [charge("Customs Fee", 1, 80.0, 80.0, "USD", "PREPAID", "Customs Fee 80.00 USD PREPAID")],
                [total("total", 80.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2005", "2026-09-26", "Destination Agent",
                    "Origin Agent", "Customs Fee", "Customs Clearance Fee", 80.0, "USD", "PP",
                    "AP", "CREATE", "NONE",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "minimal_l2",
            "file": "files/minimal_l2.xlsx",
            "bills": ["bills/BILL-MIN-2006.json"],
            "match_keywords": ["THC", "OOLU4055331281"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                None,
                "2026-09-27",
                "Destination Agent",
                "Origin Agent",
                [charge("THC", 1, 100.0, 100.0, "USD", "CC", "THC 1 100 100 USD CC")],
                [total("total", 100.0, "USD")],
                bl="OOLU4055331281",
            ),
            "expected_records": [
                expected(
                    "Debit Note", "OOLU4055331281", "2026-09-27", "Destination Agent",
                    "Origin Agent", "THC", "Terminal Handling Charge", 100.0, "USD", "CC",
                    "AP", "NOCHANGE", "L2",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "multi_currency",
            "file": "files/multi_currency.xlsx",
            "bills": [],
            "match_keywords": ["O/F", "DN-2026-2007"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2007",
                "2026-09-28",
                "Destination Agent",
                "Origin Agent",
                [
                    charge("O/F", 1, 100.0, 100.0, "USD", "PP", "O/F 1 100 100 USD PP"),
                    charge("THC", 1, 200.0, 200.0, "CNY", "CC", "THC 1 200 200 CNY CC"),
                ],
                [total("total", 100.0, "USD"), total("total", 200.0, "CNY")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2007", "2026-09-28", "Destination Agent",
                    "Origin Agent", "O/F", "Ocean Freight", 100.0, "USD", "PP",
                    "AP", "CREATE", "NONE",
                ),
                expected(
                    "Debit Note", "DN-2026-2007", "2026-09-28", "Destination Agent",
                    "Origin Agent", "THC", "Terminal Handling Charge", 200.0, "CNY", "CC",
                    "AP", "CREATE", "NONE",
                ),
            ],
        }
    )

    cases.append(
        {
            "id": "sum_mismatch",
            "file": "files/sum_mismatch.xlsx",
            "bills": [],
            "match_keywords": ["H/D Fee", "DN-2026-2008"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2008",
                "2026-09-29",
                "Destination Agent",
                "Origin Agent",
                [charge("H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID", "H/D Fee 2 20 40 USD PREPAID")],
                [total("total", 99.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2008", "2026-09-29", "Destination Agent",
                    "Origin Agent", "H/D Fee", "Handling Charge", 40.0, "USD", "PP",
                    "AP", "PENDING_REVIEW", "NONE", "SUM_MISMATCH",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "entity_unresolved",
            "file": "files/entity_unresolved.xlsx",
            "bills": [],
            "match_keywords": ["H/D Fee", "DN-2026-2009"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2009",
                "2026-09-29",
                "Unknown Co.",
                "Other Co.",
                [charge("H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID", "H/D Fee 2 20 40 USD PREPAID")],
                [total("total", 40.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2009", "2026-09-29", "Unknown Co.",
                    "Other Co.", "H/D Fee", "Handling Charge", 40.0, "USD", "PP",
                    None, "PENDING_REVIEW", "NONE", "ENTITY_UNRESOLVED",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "data_conflict",
            "file": "files/data_conflict.xlsx",
            "bills": ["bills/BILL-2010.json"],
            "match_keywords": ["H/D Fee", "DN-2026-2010"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2010",
                "2026-09-30",
                "Origin Agent",
                "Another Co.",
                [charge("H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID", "H/D Fee 2 20 40 USD PREPAID")],
                [total("total", 40.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2010", "2026-09-30", "Origin Agent",
                    "Another Co.", "H/D Fee", "Handling Charge", 40.0, "USD", "PP",
                    "AR", "PENDING_REVIEW", "L1", "DATA_CONFLICT",
                )
            ],
        }
    )

    cases.append(
        {
            "id": "ambiguous_match",
            "file": "files/ambiguous_match.xlsx",
            "bills": ["bills/BILL-2011-1.json", "bills/BILL-2011-2.json"],
            "match_keywords": ["H/D Fee", "DN-2026-2011"],
            "ocr": None,
            "extraction": extraction(
                "DEBIT NOTE",
                "DN-2026-2011",
                "2026-10-01",
                "Destination Agent",
                "Origin Agent",
                [charge("H/D Fee", 2, 20.0, 40.0, "USD", "PREPAID", "H/D Fee 2 20 40 USD PREPAID")],
                [total("total", 40.0, "USD")],
            ),
            "expected_records": [
                expected(
                    "Debit Note", "DN-2026-2011", "2026-10-01", "Destination Agent",
                    "Origin Agent", "H/D Fee", "Handling Charge", 40.0, "USD", "PP",
                    "AP", "PENDING_REVIEW", "L1", "MATCH_AMBIGUOUS",
                )
            ],
        }
    )

    write_json(ROOT / "cases.json", {"version": "2026.10.24", "cases": cases})


if __name__ == "__main__":
    build_files()
    build_bills()
    build_cases()
    print(f"generated {len(json.loads((ROOT / 'cases.json').read_text(encoding='utf-8'))['cases'])} cases")
